"""Ajoute à data/base_complete.xlsx les articles listés dans data/extra_works.txt (rejouable).

Lancer :  python scripts/extend_corpus.py   puis   python scripts/add_urls.py
"""
import copy
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend.app import config  # noqa: E402
from backend.app.corpus import load_corpus  # noqa: E402
from scripts.augment_data import (BF_FETCH_LIMIT, BF_FILTER, SELECT, _get, _is_english,  # noqa: E402
                                  authorship_names, bf_institutions, clean, fetch_all, norm_title,
                                  rebuild_abstract)

EXTRA_WORKS = config.DATA_DIR / "extra_works.txt"
_DASHES = re.compile("[" + chr(0x2014) + chr(0x2013) + "]")


def read_work_ids(path: Path) -> list[str]:
    return [s for s in (l.strip() for l in path.read_text(encoding="utf-8").splitlines()) if s and not s.startswith("#")]


def load_works(wanted: list[str]) -> dict[str, dict]:
    """Notices depuis le cache des pages BF, sinon un appel works/<id> (mis en cache disque)."""
    found = {w["id"].rsplit("/", 1)[1]: w for w in fetch_all(BF_FILTER, limit=BF_FETCH_LIMIT)
             if w.get("id", "").rsplit("/", 1)[1] in wanted}
    for wid in wanted:
        if wid not in found:
            found[wid] = _get(f"works/{wid}", {"select": SELECT})
    return found


def build_row(work: dict, new_id: str) -> list:
    title = _DASHES.sub("-", clean(work.get("title")))
    abstract = _DASHES.sub("-", clean(rebuild_abstract(work.get("abstract_inverted_index"))))
    unis = bf_institutions(work)
    if not unis:
        raise ValueError(f"{new_id} : aucune institution du Burkina Faso dans la notice")
    if not 60 <= len(abstract.split()) <= 450:
        raise ValueError(f"{new_id} : résumé de {len(abstract.split())} mots (attendu 60 à 450)")
    if not (_is_english(title) and _is_english(abstract)):
        raise ValueError(f"{new_id} : titre ou résumé non anglais")
    authors = "; ".join(_DASHES.sub("-", a) for a in authorship_names(work))
    return [new_id, title, abstract, authors, work.get("publication_year"), "; ".join(unis), ""]


def append_rows(path: Path, rows: list[list]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    model = [c for c in ws[ws.max_row]]
    for r in rows:
        ws.append(r)
        for src, dst in zip(model, ws[ws.max_row]):
            dst.font, dst.alignment = copy.copy(src.font), copy.copy(src.alignment)
    wb.save(path)


def main() -> int:
    docs = load_corpus(config.CORPUS_EXCEL)
    seen = {norm_title(d.title) for d in docs}
    wanted = read_work_ids(EXTRA_WORKS)
    works = load_works(wanted)
    links = json.loads(config.DOI_JSON.read_text(encoding="utf-8")) if config.DOI_JSON.exists() else {}
    rows, next_n = [], len(docs) + 1
    for wid in wanted:
        w = works.get(wid)
        if not w:
            print(f"{wid} : notice introuvable, ignorée")
            continue
        if norm_title(clean(w.get("title"))) in seen:
            print(f"{wid} : déjà dans la base, ignorée")
            continue
        new_id = f"Document_{next_n:02d}"
        row = build_row(w, new_id)
        rows.append(row)
        links[new_id] = w.get("doi") or w.get("id")
        seen.add(norm_title(row[1]))
        next_n += 1
    if rows:
        append_rows(config.CORPUS_EXCEL, rows)
        config.DOI_JSON.write_text(json.dumps(links, indent=1), encoding="utf-8")
    print(f"{len(rows)} article(s) ajouté(s) ; la base compte {len(docs) + len(rows)} documents.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
