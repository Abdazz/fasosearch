"""Remplit la colonne URL de data/base_complete.xlsx : page de l'article chez l'éditeur, jamais OpenAlex.

À lancer une fois, avec accès réseau :  python scripts/add_urls.py
Les réponses OpenAlex sont mises en cache dans data/openalex_cache/ : une exécution
interrompue (quota épuisé) reprend sans refaire les appels déjà faits.
"""
import copy
import json
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend.app import config  # noqa: E402
from backend.app.corpus import load_corpus  # noqa: E402
from scripts.augment_data import QuotaExceeded, _get, _significant_words, norm_title  # noqa: E402

URL_SELECT = "id,doi,title,primary_location,locations"
DOI_PREFIX = "https://doi.org/"


def _acceptable(url) -> bool:
    if not isinstance(url, str) or not url.startswith("https://"):
        return False
    host = (urlparse(url).hostname or "").lower()
    return host != "openalex.org" and not host.endswith(".openalex.org")


def choose_url(work: dict | None) -> str:
    """Page éditeur de la notice, sinon DOI, sinon autre page https hors OpenAlex, sinon ""."""
    if not work:
        return ""
    primary = (work.get("primary_location") or {}).get("landing_page_url")
    if _acceptable(primary):
        return primary
    doi = work.get("doi") or ""
    if doi.startswith(DOI_PREFIX):
        return doi
    for loc in work.get("locations") or []:
        url = (loc or {}).get("landing_page_url")
        if _acceptable(url):
            return url
    return ""


def resolve_url(link: str | None, work: dict | None) -> str:
    """URL retenue ; le DOI déjà connu (doi.json) sert de repli si la notice ne donne rien."""
    return choose_url(work) or (link if link and link.startswith(DOI_PREFIX) else "")


def fetch_work(link: str | None, title: str) -> dict | None:
    if link and link.startswith(DOI_PREFIX):
        return _get(f"works/doi:{link[len(DOI_PREFIX):]}", {"select": URL_SELECT}) or None
    if link and "openalex.org/W" in link:
        return _get(f"works/{link.rsplit('/', 1)[1]}", {"select": URL_SELECT}) or None
    data = _get("works", {"filter": f"title.search:{_significant_words(title)}", "per-page": 5,
                          "select": URL_SELECT})
    return next((w for w in data.get("results", []) if norm_title(w.get("title") or "") == norm_title(title)), None)


def write_url_column(path: Path, urls: dict[str, str]) -> None:
    """Ajoute (ou remplace) la colonne URL juste après University, sans toucher aux autres cellules."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    if "URL" in header:
        col = header.index("URL") + 1
    else:
        col = header.index("University") + 2
        if col <= ws.max_column:
            ws.insert_cols(col)
        uni = ws.cell(row=1, column=col - 1)
        cell = ws.cell(row=1, column=col, value="URL")
        cell.font, cell.fill, cell.alignment = copy.copy(uni.font), copy.copy(uni.fill), copy.copy(uni.alignment)
        ws.column_dimensions[get_column_letter(col)].width = 60
    id_col = header.index("ID_document") + 1
    for row in range(2, ws.max_row + 1):
        doc_id = ws.cell(row=row, column=id_col).value
        if doc_id:
            ws.cell(row=row, column=col, value=urls.get(str(doc_id).strip()) or None)
    wb.save(path)


def main() -> int:
    docs = load_corpus(config.CORPUS_EXCEL)
    links = json.loads(config.DOI_JSON.read_text(encoding="utf-8")) if config.DOI_JSON.exists() else {}
    urls: dict[str, str] = {}
    try:
        for d in docs:
            link = links.get(d.id)
            urls[d.id] = resolve_url(link, fetch_work(link, d.title))
    except QuotaExceeded as e:
        print(f"Quota OpenAlex épuisé : relancer dans {e.retry_after_seconds} s (le cache conserve les réponses déjà reçues). "
              "Excel non modifié.")
        return 1
    write_url_column(config.CORPUS_EXCEL, urls)
    missing = [i for i, u in urls.items() if not u]
    print(f"{len(urls) - len(missing)} URL sur {len(urls)} documents.")
    if missing:
        print("Sans URL : " + ", ".join(missing))
    return 0


if __name__ == "__main__":
    sys.exit(main())
