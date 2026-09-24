"""Représentation d'un document et lecture de la base Excel."""
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl

COLUMNS = ["ID_document", "Title", "Abstract", "Authors", "Year", "University"]

_BARE_DASH_RE = re.compile("[\u2014\u2013]")


def normalize_dashes(s: str) -> str:
    """Normalise les tirets cadratins et demi-cadratins en tiret simple.

    - Le tiret cadratin ou demi-cadratin entouré d'espaces (U+2014 ou U+2013) devient " - ".
    - Le tiret cadratin ou demi-cadratin sans espaces (ex. 2010 avec U+2013 2020) devient "-".
    """
    return _BARE_DASH_RE.sub("-", s)


@dataclass(frozen=True)
class Document:
    id: str
    title: str
    abstract: str
    authors: str
    year: int | None
    university: str

    @property
    def text(self) -> str:
        """Texte indexé : titre + résumé."""
        return f"{self.title}. {self.abstract}"

    def to_dict(self) -> dict:
        return asdict(self)


def _year(v) -> int | None:
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def load_corpus(path: Path) -> list[Document]:
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    col = {name: header.index(name) for name in COLUMNS if name in header}
    docs = []
    for r in rows[1:]:
        if not r or not r[col["ID_document"]]:
            continue
        get = lambda name: r[col[name]] if name in col and r[col[name]] is not None else ""  # noqa: E731
        docs.append(Document(
            id=str(get("ID_document")).strip(),
            title=normalize_dashes(str(get("Title")).strip()),
            abstract=normalize_dashes(str(get("Abstract")).strip()),
            authors=normalize_dashes(str(get("Authors")).strip()),
            year=_year(get("Year")),
            university=normalize_dashes(str(get("University")).strip()),
        ))
    return docs
