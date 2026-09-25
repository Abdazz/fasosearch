import openpyxl
import pytest

from scripts import extend_corpus as E


def work(title="A Blockchain Protocol for Mobile Money", country="BF", n_words=80, authors=("Kodjo Agbezoutsi", "Burkina Faso")):
    words = ("blockchain " * n_words).split()
    words[3] = "security " + chr(0x2014) + " layer"
    return {
        "id": "https://openalex.org/W1", "doi": "https://doi.org/10.1/x", "title": title, "publication_year": 2019,
        "abstract_inverted_index": {w: [i] for i, w in enumerate(words)} | {"blockchain": [i for i, w in enumerate(words) if w == "blockchain"]},
        "authorships": [{"author": {"display_name": a}, "institutions": [{"display_name": "Nazi Boni University", "country_code": country}]} for a in authors],
    }


def test_read_work_ids(tmp_path):
    p = tmp_path / "ids.txt"
    p.write_text("# commentaire\nW1\n\n  W2  \n", encoding="utf-8")
    assert E.read_work_ids(p) == ["W1", "W2"]


def test_build_row_cleans_and_filters():
    row = E.build_row(work(), "Document_91")
    assert row[0] == "Document_91" and row[4] == 2019
    assert row[3] == "Kodjo Agbezoutsi"
    assert row[5] == "Nazi Boni University"
    assert chr(0x2014) not in row[2] and chr(0x2013) not in row[2]
    assert row[6] == ""


def test_build_row_rejects_non_bf_and_bad_length():
    with pytest.raises(ValueError, match="Burkina"):
        E.build_row(work(country="FR"), "Document_91")
    with pytest.raises(ValueError, match="mots"):
        E.build_row(work(n_words=20), "Document_91")


def _book(path):
    wb = openpyxl.Workbook()
    wb.active.append(["ID_document", "Title", "Abstract", "Authors", "Year", "University", "URL"])
    wb.active.append(["Document_01", "Old title", "A", "X Y", 2020, "U", "https://doi.org/1"])
    wb.save(path)


def test_extend_is_idempotent(tmp_path, monkeypatch):
    p = tmp_path / "b.xlsx"
    _book(p)
    ids = tmp_path / "ids.txt"
    ids.write_text("W1\n", encoding="utf-8")
    doi = tmp_path / "doi.json"
    doi.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(E.config, "CORPUS_EXCEL", p)
    monkeypatch.setattr(E.config, "DOI_JSON", doi)
    monkeypatch.setattr(E, "EXTRA_WORKS", ids)
    monkeypatch.setattr(E, "load_works", lambda wanted: {"W1": work()})
    assert E.main() == 0
    assert E.main() == 0
    rows = list(openpyxl.load_workbook(p).active.iter_rows(values_only=True))
    assert [r[0] for r in rows[1:]] == ["Document_01", "Document_02"]
    assert rows[1] == ("Document_01", "Old title", "A", "X Y", 2020, "U", "https://doi.org/1")
    assert '"Document_02": "https://doi.org/10.1/x"' in doi.read_text(encoding="utf-8")


def test_duplicate_title_is_refused(tmp_path, monkeypatch):
    p = tmp_path / "b.xlsx"
    _book(p)
    ids = tmp_path / "ids.txt"
    ids.write_text("W1\n", encoding="utf-8")
    doi = tmp_path / "doi.json"
    doi.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(E.config, "CORPUS_EXCEL", p)
    monkeypatch.setattr(E.config, "DOI_JSON", doi)
    monkeypatch.setattr(E, "EXTRA_WORKS", ids)
    monkeypatch.setattr(E, "load_works", lambda wanted: {"W1": work(title="Old title")})
    assert E.main() == 0
    assert len(list(openpyxl.load_workbook(p).active.iter_rows())) == 2
