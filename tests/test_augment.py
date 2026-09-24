import json

import openpyxl
import pytest
import requests

from backend.app import config
from backend.app.corpus import COLUMNS
from scripts.augment_data import (CS_LEXICON, PER_SUBFIELD, SEARCH_RETRY_MAX_WAIT,
                                  STRONG_LEXICON, WEAK_LEXICON, QuotaExceeded, bf_institutions,
                                  clean, is_computer_science, is_excluded, load_exclusions,
                                  next_ids, norm_title, rebuild_abstract, select_diverse,
                                  write_excel,
                                  _get, _search_paced, _significant_words)


def test_rebuild_abstract_orders_words():
    inv = {"world": [1], "hello": [0], "again": [3], "hello_": [2]}
    assert rebuild_abstract(inv) == "hello world hello_ again"
    assert rebuild_abstract(None) == ""


def test_norm_title_ignores_case_accents_punctuation():
    assert norm_title("Réseaux: l'Étude!") == norm_title("reseaux letude")


def test_clean_strips_tags_and_abstract_prefix():
    assert clean("<p>Abstract: Hello   <b>world</b></p>") == "Hello world"


def test_bf_institutions_keeps_only_burkina():
    work = {"authorships": [
        {"institutions": [{"display_name": "Université Norbert Zongo", "country_code": "BF"}]},
        {"institutions": [{"display_name": "Sorbonne", "country_code": "FR"},
                          {"display_name": "Université Norbert Zongo", "country_code": "BF"}]},
        {"institutions": [{"display_name": "Université Joseph Ki-Zerbo", "country_code": "BF"}]},
    ]}
    assert bf_institutions(work) == ["Université Norbert Zongo", "Université Joseph Ki-Zerbo"]


def test_select_diverse_caps_subfields_and_prefers_recent():
    cands = [{"id": i, "year": 2000 + i, "subfield": "A" if i < 6 else "B"} for i in range(10)]
    chosen = select_diverse(cands, n=5, per_subfield=2)
    assert [c["id"] for c in chosen] == [9, 8, 5, 4]  # 2 de B, 2 de A ; plus assez pour 5
    assert len(select_diverse(cands, n=3, per_subfield=5)) == 3


def test_next_ids_continue_numbering():
    assert next_ids(31, 3) == ["Document_31", "Document_32", "Document_33"]
    assert next_ids(100, 1) == ["Document_100"]


def test_per_subfield_is_15():
    assert PER_SUBFIELD == 15


def test_cs_lexicon_has_at_least_60_terms():
    assert len(CS_LEXICON) >= 60
    assert all(t == t.lower() for t in CS_LEXICON)
    assert set(CS_LEXICON) == set(STRONG_LEXICON) | set(WEAK_LEXICON)


def test_strong_and_weak_lexicons_are_disjoint():
    assert set(STRONG_LEXICON) & set(WEAK_LEXICON) == set()


def test_is_computer_science_accepts_clear_cs_example():
    title = "A Deep Learning Approach for Malware Classification"
    abstract = ("We propose a machine learning based classifier using a convolutional "
                "neural network architecture to analyze software behavior.")
    assert is_computer_science(title, abstract) is True


def test_is_computer_science_rejects_health_study_with_only_generic_terms():
    title = "Statistical Analysis of Patient Health Data in Rural Clinics"
    abstract = ("This study uses a statistical model to analyze patient data collected "
                "from rural health clinics, examining trends in disease prevalence.")
    assert is_computer_science(title, abstract) is False


def test_is_computer_science_rejects_economics_paper_mentioning_digital_platform():
    title = "Economic Impact of Digital Platform Adoption on Rural Markets"
    abstract = ("This paper studies how adoption of a digital platform affects income and "
                "market participation among smallholder farmers, using survey data and a "
                "regression model.")
    assert is_computer_science(title, abstract) is False


def test_is_computer_science_needs_at_least_one_strong_term_in_title_when_only_two_total():
    # 2 termes STRONG au total mais aucun dans le titre -> refusé (il en faudrait 3).
    title = "Findings From a Field Study in West Africa"
    abstract = "The team used a database and a classifier to organize survey responses."
    assert is_computer_science(title, abstract) is False


def test_is_computer_science_accepts_three_strong_terms_even_without_title_match():
    title = "Findings From a Field Study in West Africa"
    abstract = "The team used a database, a classifier and a neural network to analyze survey responses."
    assert is_computer_science(title, abstract) is True


def test_is_computer_science_is_case_insensitive():
    title = "MACHINE LEARNING for Intrusion Detection"
    abstract = "We use a CLASSIFIER and a NEURAL NETWORK."
    assert is_computer_science(title, abstract) is True


# ---------------------------------------------------------------- exclusion manuelle
def test_load_exclusions_ignores_blank_and_comment_lines(tmp_path):
    p = tmp_path / "exclusions.txt"
    p.write_text(f"# commentaire\n\n{norm_title('Some Title Here')}\nW12345\n", encoding="utf-8")
    assert load_exclusions(p) == {norm_title("Some Title Here"), "W12345"}


def test_load_exclusions_missing_file_returns_empty_set(tmp_path):
    assert load_exclusions(tmp_path / "missing.txt") == set()


def test_is_excluded_matches_by_normalized_title_or_id():
    work = {"id": "https://openalex.org/W999"}
    assert is_excluded(work, "Some Title!", {norm_title("some title")})
    assert is_excluded(work, "Other", {"W999"})
    assert is_excluded(work, "Other", {"https://openalex.org/W999"})
    assert not is_excluded(work, "Other", {"nope"})


# ---------------------------------------------------------------- cache disque / quota
class _FakeResponse:
    def __init__(self, status_code=200, json_data=None, headers=None):
        self.status_code = status_code
        self._json = json_data or {}
        self.headers = headers or {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"status {self.status_code}")

    def json(self):
        return self._json


@pytest.fixture
def cache_dir(tmp_path, monkeypatch):
    d = tmp_path / "openalex_cache"
    monkeypatch.setattr(config, "OPENALEX_CACHE_DIR", d)
    return d


def test_get_reads_from_cache_without_network_call(cache_dir, monkeypatch):
    calls = []
    monkeypatch.setattr(requests, "get", lambda *a, **k: calls.append(1) or _FakeResponse())
    first = _get("works", {"filter": "x", "per-page": 1})
    assert len(calls) == 1
    second = _get("works", {"filter": "x", "per-page": 1})
    assert len(calls) == 1  # pas de second appel réseau : lu depuis le cache
    assert first == second


def test_get_writes_cache_file_on_success(cache_dir, monkeypatch):
    monkeypatch.setattr(requests, "get", lambda *a, **k: _FakeResponse(json_data={"ok": True}))
    _get("works", {"filter": "y"})
    files = list(cache_dir.glob("*.json"))
    assert len(files) == 1
    assert json.loads(files[0].read_text()) == {"ok": True}


def test_get_raises_quota_exceeded_on_429_without_retry_loop(cache_dir, monkeypatch):
    calls = []

    def fake_get(*a, **k):
        calls.append(1)
        return _FakeResponse(status_code=429, headers={"Retry-After": "120"})

    monkeypatch.setattr(requests, "get", fake_get)
    monkeypatch.setattr("time.sleep", lambda *_: None)
    with pytest.raises(QuotaExceeded) as exc_info:
        _get("works", {"filter": "z"})
    assert exc_info.value.retry_after_seconds == 120
    assert len(calls) == 1  # aucune nouvelle tentative sur 429


def test_get_retries_three_times_on_other_network_errors(cache_dir, monkeypatch):
    calls = []

    def fake_get(*a, **k):
        calls.append(1)
        raise requests.ConnectionError("boom")

    monkeypatch.setattr(requests, "get", fake_get)
    monkeypatch.setattr("time.sleep", lambda *_: None)
    result = _get("works", {"filter": "w"})
    assert result == {}
    assert len(calls) == 4  # 1 essai + 3 nouvelles tentatives


# ---------------------------------------------------------------- affiliations d'origine
def test_significant_words_keeps_first_n_words_over_two_letters():
    title = "A Relevant Feature Identification Approach to Detect APTs in HTTPS Traffic Today"
    assert _significant_words(title, 8) == "Relevant Feature Identification Approach Detect APTs HTTPS Traffic"


def test_significant_words_handles_empty_title():
    assert _significant_words("") == ""
    assert _significant_words(None) == ""


def test_search_paced_sleeps_at_least_min_interval_between_uncached_calls(cache_dir, monkeypatch):
    monkeypatch.setattr(requests, "get", lambda *a, **k: _FakeResponse(json_data={"results": []}))
    sleeps = []
    monkeypatch.setattr("time.sleep", lambda s: sleeps.append(s))
    import scripts.augment_data as aug
    aug._last_search_call[0] = 0.0
    monkeypatch.setattr("time.time", lambda: 1000.0)
    _search_paced({"search": "first title", "per-page": 5})
    aug._last_search_call[0] = 1000.5  # moins de 2s depuis le dernier appel réseau
    monkeypatch.setattr("time.time", lambda: 1000.5)
    _search_paced({"search": "second title", "per-page": 5})
    assert any(s >= 1.4 for s in sleeps)  # a bien attendu pour respecter les 2s


def test_search_paced_retries_once_on_short_429_then_succeeds(cache_dir, monkeypatch):
    calls = []

    def fake_get(*a, **k):
        calls.append(1)
        if len(calls) == 1:
            return _FakeResponse(status_code=429, headers={"Retry-After": "5"})
        return _FakeResponse(json_data={"results": [{"title": "ok"}]})

    monkeypatch.setattr(requests, "get", fake_get)
    monkeypatch.setattr("time.sleep", lambda *_: None)
    data = _search_paced({"search": "some title", "per-page": 5})
    assert data == {"results": [{"title": "ok"}]}
    assert len(calls) == 2  # 1 essai (429) + 1 retentative (succès)


def test_write_excel_header_has_url_and_widths_cover_column_g(tmp_path):
    # COLUMNS compte 7 noms (dont URL) alors que les lignes écrites par _run() n'en ont que 6
    # (ce script ne connaît pas la page éditeur) : la colonne URL doit rester vide, mais
    # l'en-tête et la largeur de colonne (G) doivent quand même la couvrir, pour que
    # scripts/add_urls.py la retrouve et la remplisse ensuite.
    path = tmp_path / "out.xlsx"
    write_excel([["Document_01", "Titre", "Résumé", "Auteur", 2020, "Université X"]], path)

    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS == [
        "ID_document", "Title", "Abstract", "Authors", "Year", "University", "URL"]
    assert ws.column_dimensions["G"].width == 60
    assert ws.cell(row=2, column=7).value is None  # URL laissée vide par ce script


def test_search_paced_gives_up_cleanly_on_long_429(cache_dir, monkeypatch):
    monkeypatch.setattr(requests, "get",
                         lambda *a, **k: _FakeResponse(status_code=429, headers={"Retry-After": "999"}))
    monkeypatch.setattr("time.sleep", lambda *_: None)
    with pytest.raises(QuotaExceeded) as exc_info:
        _search_paced({"search": "some title", "per-page": 5})
    assert exc_info.value.retry_after_seconds == 999
    assert 999 > SEARCH_RETRY_MAX_WAIT
