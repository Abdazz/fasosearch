import json

import openpyxl
import pytest

from scripts import add_urls
from scripts.add_urls import choose_url, fetch_work, resolve_url, write_url_column


def work(primary=None, doi=None, locations=()):
    return {"doi": doi, "primary_location": {"landing_page_url": primary} if primary is not None else None,
            "locations": [{"landing_page_url": u} for u in locations]}


def test_choose_url_prefers_publisher_page():
    assert choose_url(work("https://ieeexplore.ieee.org/document/1", "https://doi.org/10.1/x")) == "https://ieeexplore.ieee.org/document/1"


def test_choose_url_rejects_openalex_and_uses_doi():
    assert choose_url(work("https://openalex.org/W1", "https://doi.org/10.1/x")) == "https://doi.org/10.1/x"


def test_choose_url_rejects_http():
    assert choose_url(work("http://example.org/a", "https://doi.org/10.1/x")) == "https://doi.org/10.1/x"


def test_choose_url_falls_back_to_locations():
    w = work(None, None, ["https://openalex.org/W2", "http://x.org", "https://arxiv.org/abs/2401.00001"])
    assert choose_url(w) == "https://arxiv.org/abs/2401.00001"


def test_choose_url_empty():
    assert choose_url(None) == "" and choose_url({}) == ""
    assert choose_url(work(None, None, ["https://api.openalex.org/works/W3"])) == ""


def test_resolve_url_falls_back_to_known_doi():
    assert resolve_url("https://doi.org/10.1/x", None) == "https://doi.org/10.1/x"
    assert resolve_url("https://openalex.org/W1", None) == ""
    assert resolve_url(None, work("https://hal.science/hal-1")) == "https://hal.science/hal-1"


def test_fetch_work_routes(monkeypatch):
    calls = []

    def fake_get(path, params):
        calls.append(path)
        if path == "works":
            return {"results": [{"title": "Other"}, {"title": "My Paper!", "doi": "https://doi.org/10.9/z"}]}
        return {"doi": "https://doi.org/10.1/x"}

    monkeypatch.setattr(add_urls, "_get", fake_get)
    assert fetch_work("https://doi.org/10.1/x", "T")["doi"] == "https://doi.org/10.1/x"
    assert fetch_work("https://openalex.org/W42", "T") is not None
    assert fetch_work(None, "My paper")["doi"] == "https://doi.org/10.9/z"
    assert fetch_work(None, "Absent title") is None
    assert calls[:2] == ["works/doi:10.1/x", "works/W42"]


def _book(path):
    wb = openpyxl.Workbook()
    wb.active.append(["ID_document", "Title", "Abstract", "Authors", "Year", "University"])
    wb.active.append(["Document_01", "T1", "A1", "X", 2020, "U1"])
    wb.active.append(["Document_02", "T2", "A2", "Y", 2021, "U2"])
    wb.save(path)


def _rows(path):
    return [list(r) for r in openpyxl.load_workbook(path).active.iter_rows(values_only=True)]


def test_write_url_column_after_university(tmp_path):
    p = tmp_path / "b.xlsx"
    _book(p)
    write_url_column(p, {"Document_01": "https://doi.org/10.1/x"})
    rows = _rows(p)
    assert rows[0] == ["ID_document", "Title", "Abstract", "Authors", "Year", "University", "URL"]
    assert rows[1] == ["Document_01", "T1", "A1", "X", 2020, "U1", "https://doi.org/10.1/x"]
    assert rows[2][:6] == ["Document_02", "T2", "A2", "Y", 2021, "U2"] and rows[2][6] in (None, "")


def test_write_url_column_is_idempotent(tmp_path):
    p = tmp_path / "b.xlsx"
    _book(p)
    write_url_column(p, {"Document_01": "https://a.org/1"})
    write_url_column(p, {"Document_01": "https://b.org/2", "Document_02": "https://c.org/3"})
    rows = _rows(p)
    assert rows[0].count("URL") == 1 and len(rows[0]) == 7
    assert rows[1][6] == "https://b.org/2" and rows[2][6] == "https://c.org/3"


def test_main_does_not_write_on_quota(tmp_path, monkeypatch):
    p = tmp_path / "b.xlsx"
    _book(p)
    doi = tmp_path / "doi.json"
    doi.write_text(json.dumps({"Document_01": "https://doi.org/10.1/x"}), encoding="utf-8")
    monkeypatch.setattr(add_urls.config, "CORPUS_EXCEL", p)
    monkeypatch.setattr(add_urls.config, "DOI_JSON", doi)

    def quota(*_):
        raise add_urls.QuotaExceeded(3600)

    monkeypatch.setattr(add_urls, "fetch_work", quota)
    before = p.read_bytes()
    assert add_urls.main() == 1
    assert p.read_bytes() == before
